import io
import re
import logging

from openpyxl import load_workbook
from openpyxl.styles import Font

from models.schemas import PatientData
from utils.paths import get_template_xlsx

logger = logging.getLogger(__name__)

TEMPLATE_PATH = get_template_xlsx()

_BOLD = Font(bold=True)


def generate_excel(data: PatientData) -> bytes:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"Plantilla no encontrada en {TEMPLATE_PATH}. "
            "Copie el archivo a backend/templates/plantilla_comite_tumores.xlsx"
        )

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb["Hoja1"]

    def put(cell: str, value, bold: bool = False):
        if value:
            ws[cell] = str(value).strip()
            if bold:
                ws[cell].font = _BOLD

    # Sección I — Identificación
    put("F4",  data.documento)

    if data.fecha_atencion:
        d, m, a = _split_date(data.fecha_atencion)
        put("J4", d)
        put("K4", m)
        put("L4", a, bold=True)        # Año en negrita

    put("P4",  data.edad)
    put("D6",  data.nombre)
    put("J6",  data.apellidos)
    put("F8",  data.medico_tratante)
    put("P8",  data.especialidad, bold=True)   # Especialidad en negrita
    put("E10", data.fecha_diagnostico)
    put("J10", data.entidad)

    # Sección II — Resumen
    put("B14", data.resumen_historia)

    # Sección III — Diagnóstico
    put("G25", data.diagnostico_clinico, bold=True)   # Dx clínico en negrita
    put("N25", data.diagnostico_patologico)

    put("C27", data.tnm_t)
    put("E27", data.tnm_n)
    put("G27", data.tnm_m)
    put("I27", data.estadificacion)

    # Imágenes
    put("D31", data.ecografia)
    put("F31", data.tac)
    put("H31", data.rmn)
    put("J31", data.pet)
    put("O31", data.mamografia)
    put("B34", data.informes_imagenes)

    # Tratamientos
    put("D38", data.cirugia_recibio)
    put("F38", data.cirugia_fecha)
    put("J38", data.cirugia_tipo)

    put("E40", data.quimioterapia_recibio)
    put("I40", data.quimioterapia_esquema)

    put("E42", data.radioterapia_recibio)
    put("G42", data.radioterapia_dosis)

    put("E44", data.hormonoterapia_recibio)
    put("G44", data.hormonoterapia_tipo)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _split_date(date_str: str) -> tuple[str, str, str]:
    """Devuelve (día, mes, año) desde varios formatos de fecha."""
    date_str = date_str.strip()

    # "2025 - Enero - 20"
    m = re.match(r"(\d{4})\s*[-–]\s*(\w+)\s*[-–]\s*(\d{1,2})", date_str, re.IGNORECASE)
    if m:
        return m.group(3).zfill(2), _month_to_num(m.group(2)), m.group(1)

    # ISO YYYY-MM-DD
    m = re.match(r"(\d{4})[\/\-\.](\d{1,2})[\/\-\.](\d{1,2})", date_str)
    if m:
        return m.group(3).zfill(2), m.group(2).zfill(2), m.group(1)

    # DD/MM/YYYY
    m = re.match(r"(\d{1,2})[\/\-\.](\d{1,2})[\/\-\.](\d{2,4})", date_str)
    if m:
        return m.group(1).zfill(2), m.group(2).zfill(2), m.group(3)

    return date_str, "", ""


def _month_to_num(name: str) -> str:
    return {
        "enero": "01", "febrero": "02", "marzo": "03", "abril": "04",
        "mayo": "05", "junio": "06", "julio": "07", "agosto": "08",
        "septiembre": "09", "octubre": "10", "noviembre": "11", "diciembre": "12",
        "january": "01", "february": "02", "march": "03", "april": "04",
        "may": "05", "june": "06", "july": "07", "august": "08",
        "september": "09", "october": "10", "november": "11", "december": "12",
    }.get(name.lower(), name)
